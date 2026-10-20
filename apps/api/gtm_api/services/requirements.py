"""Requirement-to-proof: extract customer requirements and answer each from evidence.

Pipeline per requirement: retrieve passages -> LLM drafts verdict/answer/conditions ->
the claim-level citation gate re-checks the answer -> verdict is downgraded when the
evidence does not carry it. Roadmap wording in a cited passage forces ``roadmap`` so a
planned capability is never presented as available.
"""

import csv
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from gtm_api.models import PROOF_VERDICTS
from gtm_api.services.citation_gate import (
    SYSTEM_PROMPT_GROUNDED,
    build_context_from_results,
    verify_grounding,
)
from gtm_api.services.llm import get_chat_model
from gtm_api.services.loaders.file import FileLoader
from gtm_api.services.vector_store import SearchResult

MAX_REQUIREMENTS = 100
_WINDOW_CHARS = 12000
_MAX_WINDOWS = 8  # ~96k characters read per document
_AGENT = "solution_architect"  # reuse existing per-agent model routing
_ROADMAP = re.compile(
    r"\b(roadmap|planned|upcoming|coming soon|future release|not yet (?:available|supported)|"
    r"under development|will be (?:added|supported|available))\b",
    re.IGNORECASE,
)
_LIST_PREFIX = re.compile(r"^\s*(?:[-*•]|\(?\d+[.)]|\(?[a-z][.)]|req[-\s]?\d+[:.)]?)\s+", re.IGNORECASE)

EXTRACT_PROMPT = """Extract the individual customer requirements from the text below.
Return ONLY a JSON array of strings. Each string is one atomic, self-contained requirement
(split compound questions like "support A, B and C" into separate items). Keep the
customer's wording; do not add requirements that are not in the text.

Text:
{text}
"""

ASSESS_PROMPT = """Customer requirement:
{requirement}

Product documentation excerpts:
{context}

Decide whether the product meets the requirement, using ONLY the excerpts.
Return ONLY JSON with these keys:
- "verdict": "supported" (fully documented), "partial" (documented with limits or only
  some of it), "roadmap" (planned, not available today), "gap" (docs say it is not
  supported), or "unknown" (excerpts do not say).
- "answer": 1-3 sentences stating what is documented. Cite excerpts as [Source N].
  Do not claim anything the excerpts do not say.
- "conditions": limits, prerequisites or caveats, or "".
- "validation_test": one concrete test the customer could run to confirm, or "".
"""


@dataclass
class ProofDraft:
    verdict: str
    answer: str
    conditions: str = ""
    validation_test: str = ""
    citations: list[dict] = field(default_factory=list)
    claims: list[dict] = field(default_factory=list)
    confidence: float = 0.0


_CUE = re.compile(
    r"\b(must|shall|should|required?|requirements?|need|needs|support|supports|provide|provides|"
    r"ability|able|capable|can|could|do you|does|will|ensure|include|includes|compatible|integrate)\b|\?",
    re.IGNORECASE,
)


def split_requirements_heuristic(text: str) -> list[str]:
    """Deterministic fallback. List items are always requirements; free-form lines count only
    when they carry a requirement cue (must/shall/support/?...), so background prose is skipped."""
    items: list[str] = []
    for raw in text.splitlines():
        is_list_item = bool(_LIST_PREFIX.match(raw))
        line = _LIST_PREFIX.sub("", raw).strip()
        if len(line.split()) < 4:
            continue
        for sentence in re.split(r"(?<=[.?!])\s+(?=[A-Z])", line):
            sentence = sentence.strip()
            if len(sentence.split()) >= 4 and (is_list_item or _CUE.search(sentence)):
                items.append(sentence)
    return _dedupe(items)


def _dedupe(items: list[str]) -> list[str]:
    seen: set[str] = set()
    out = []
    for item in items:
        key = re.sub(r"\W+", " ", item.lower()).strip()
        if key and key not in seen:
            seen.add(key)
            out.append(item.strip())
    return out


def _parse_json(content: str):
    if "```" in content:
        parts = content.split("```")
        content = parts[1].removeprefix("json") if len(parts) > 1 else content
    return json.loads(content.strip())


def _windows(text: str) -> tuple[list[str], bool]:
    """Split ``text`` on line boundaries into <= _WINDOW_CHARS windows; at most _MAX_WINDOWS.
    Returns (windows, truncated)."""
    windows: list[str] = []
    current: list[str] = []
    size = 0
    for line in text.splitlines():
        if size + len(line) + 1 > _WINDOW_CHARS and current:
            windows.append("\n".join(current))
            current, size = [], 0
        current.append(line[:_WINDOW_CHARS])
        size += len(line) + 1
    if current:
        windows.append("\n".join(current))
    return windows[:_MAX_WINDOWS], len(windows) > _MAX_WINDOWS


async def _extract_window(window: str) -> list[str]:
    try:
        llm = get_chat_model(_AGENT, temperature=0.0)
        response = await llm.ainvoke([{"role": "user", "content": EXTRACT_PROMPT.format(text=window)}])
        parsed = _parse_json(response.content)
        items = [str(i).strip() for i in parsed if isinstance(i, str) and len(i.split()) >= 3]
        if items:
            return items
    except Exception:
        pass
    return split_requirements_heuristic(window)


async def extract_requirements_with_info(text: str) -> tuple[list[str], bool]:
    """LLM extraction per window (splits compound requirements), heuristic fallback per window.
    Returns (requirements, truncated) where truncated means the text exceeded what we read."""
    text = text.strip()
    if not text:
        return [], False
    windows, truncated = _windows(text)
    items: list[str] = []
    for window in windows:
        items.extend(await _extract_window(window))
    unique = _dedupe(items)
    return unique[:MAX_REQUIREMENTS], truncated or len(unique) > MAX_REQUIREMENTS


async def extract_requirements(text: str) -> list[str]:
    return (await extract_requirements_with_info(text))[0]


MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # RFPs are documents, not bulk source uploads
UPLOAD_EXTENSIONS = (".txt", ".md", ".markdown", ".pdf", ".docx", ".pptx", ".csv", ".xlsx")


def _spreadsheet_lines(data: bytes, ext: str) -> str:
    """One line per row: the longest cell, which in a questionnaire is the requirement text
    (the other columns are ids, categories and response slots)."""
    rows: list[list[str]] = []
    if ext == ".csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8", errors="replace"))))
    else:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                rows.append([str(c) for c in row if c is not None])
        wb.close()
    lines = []
    for row in rows:
        cells = [c.strip() for c in row if c and c.strip()]
        if cells:
            lines.append(max(cells, key=len))
    return "\n".join(lines)


def extract_text_from_upload(filename: str, content_type: str | None, data: bytes) -> str:
    """Plain text from an uploaded RFP/questionnaire. Raises ValueError for unsupported types."""
    ext = Path(filename).suffix.lower()
    if ext not in UPLOAD_EXTENSIONS:
        raise ValueError(f"Unsupported file type '{ext or 'unknown'}'. Use: {', '.join(UPLOAD_EXTENSIONS)}")
    if ext in (".csv", ".xlsx"):
        return _spreadsheet_lines(data, ext)
    return FileLoader()._extract(data, ext, (content_type or "").lower())


def finalize_verdict(
    raw: dict, answer: str, results: list[SearchResult]
) -> ProofDraft:
    """Turn the model's draft into a ProofDraft, enforcing evidence rules."""
    verdict = str(raw.get("verdict", "unknown")).lower()
    if verdict not in PROOF_VERDICTS:
        verdict = "unknown"

    grounding = verify_grounding(answer, results) if results else None
    citations = [c.model_dump() for c in grounding.citations] if grounding else []
    claims = (
        [{"claim": c.claim, "status": c.status, "chunk_id": c.chunk_id} for c in grounding.claims]
        if grounding
        else []
    )

    if not results or grounding is None:
        verdict, answer = "unknown", "No documentation was found for this requirement."
    elif verdict in ("supported", "partial", "roadmap", "gap") and not grounding.grounded:
        # The evidence does not carry the claim: do not assert it to a customer.
        verdict = "unknown"
    elif verdict in ("supported", "partial"):
        cited_ids = {c["chunk_id"] for c in citations}
        if any(_ROADMAP.search(r.content) for r in results if r.chunk_id in cited_ids):
            verdict = "roadmap"

    return ProofDraft(
        verdict=verdict,
        answer=answer,
        conditions=str(raw.get("conditions") or ""),
        validation_test=str(raw.get("validation_test") or ""),
        citations=citations,
        claims=claims,
        confidence=grounding.confidence if grounding and verdict != "unknown" else 0.0,
    )


async def assess_requirement(requirement: str, results: list[SearchResult]) -> ProofDraft:
    if not results:
        return finalize_verdict({}, "", [])
    llm = get_chat_model(_AGENT, temperature=0.1)
    response = await llm.ainvoke(
        [
            {"role": "system", "content": SYSTEM_PROMPT_GROUNDED},
            {
                "role": "user",
                "content": ASSESS_PROMPT.format(
                    requirement=requirement, context=build_context_from_results(results)
                ),
            },
        ]
    )
    try:
        raw = _parse_json(response.content)
        if not isinstance(raw, dict):
            raise TypeError("not an object")
    except (TypeError, ValueError, IndexError):
        raw = {"verdict": "unknown", "answer": response.content}
    return finalize_verdict(raw, str(raw.get("answer", "")), results)


_SECTIONS = (
    ("supported", "Supported"),
    ("partial", "Supported with conditions"),
    ("roadmap", "Roadmap (not available today)"),
    ("gap", "Not supported"),
    ("unknown", "Open questions"),
)


def build_response_document(items: list[dict]) -> str:
    """Customer-ready markdown from reviewer-approved proofs.

    ``items``: [{"requirement", "verdict", "answer", "conditions", "validation_test"}].
    Deterministic on purpose: nothing is added that a reviewer did not approve.
    """
    lines = ["# Technical response to your requirements", ""]
    for verdict, heading in _SECTIONS:
        group = [i for i in items if i["verdict"] == verdict]
        if not group:
            continue
        lines += [f"## {heading}", ""]
        for item in group:
            lines.append(f"**{item['requirement']}**")
            if verdict == "unknown":
                lines.append("We are confirming this and will follow up.")
            else:
                lines.append(item["answer"])
            if item.get("conditions"):
                lines.append(f"*Conditions:* {item['conditions']}")
            if item.get("validation_test") and verdict != "gap":
                lines.append(f"*Suggested validation:* {item['validation_test']}")
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


