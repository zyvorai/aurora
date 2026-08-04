"""CSV and XLSX spreadsheet loader."""

from __future__ import annotations

import csv
import io
from pathlib import Path

from sqlalchemy.ext.asyncio import AsyncSession

from gtm_api.models import Source
from gtm_api.services.loaders.base import page_from_text
from gtm_api.services.storage import storage_service


class SpreadsheetLoader:
    async def load(self, source: Source, db: AsyncSession) -> list:
        if not source.storage_key:
            raise ValueError("Spreadsheet source missing storage_key")

        data = storage_service.get_object(source.storage_key)
        filename = source.display_name or Path(source.storage_key).name
        ext = Path(filename).suffix.lower()
        url = source.url or f"file://{filename}"

        if ext == ".csv" or (source.mime_type or "").startswith("text/csv"):
            text = self._csv_to_markdown(data)
            return [page_from_text(url, filename, text)]

        if ext in (".xlsx", ".xls") or "spreadsheet" in (source.mime_type or ""):
            return self._xlsx_pages(data, url, filename)

        raise ValueError(f"Unsupported spreadsheet type: {ext}")

    def _csv_to_markdown(self, data: bytes) -> str:
        text = data.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        if not rows:
            return ""
        header = rows[0]
        lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * len(header)) + " |"]
        for row in rows[1:]:
            padded = row + [""] * (len(header) - len(row))
            lines.append("| " + " | ".join(padded[: len(header)]) + " |")
        return "\n".join(lines)

    def _xlsx_pages(self, data: bytes, url: str, filename: str) -> list:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        pages = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append([str(c) if c is not None else "" for c in row])
            if not rows:
                continue
            header = rows[0]
            lines = [f"# Sheet: {sheet_name}", ""]
            lines.append("| " + " | ".join(header) + " |")
            lines.append("| " + " | ".join(["---"] * len(header)) + " |")
            for row in rows[1:]:
                padded = row + [""] * (len(header) - len(row))
                lines.append("| " + " | ".join(padded[: len(header)]) + " |")
            pages.append(
                page_from_text(f"{url}#{sheet_name}", f"{filename} — {sheet_name}", "\n".join(lines))
            )
        wb.close()
        return pages or [page_from_text(url, filename, "(empty spreadsheet)")]
