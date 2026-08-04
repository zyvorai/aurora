"""Extract text from uploaded files stored in MinIO."""

from __future__ import annotations

import io
from pathlib import Path

from sqlalchemy.ext.asyncio import AsyncSession

from gtm_api.models import Source
from gtm_api.services.loaders.base import page_from_text
from gtm_api.services.storage import storage_service


class FileLoader:
    async def load(self, source: Source, db: AsyncSession) -> list:
        if not source.storage_key:
            raise ValueError("File source missing storage_key")

        data = storage_service.get_object(source.storage_key)
        filename = source.display_name or Path(source.storage_key).name
        mime = (source.mime_type or "").lower()
        ext = Path(filename).suffix.lower()

        text = self._extract(data, ext, mime)
        url = source.url or f"file://{filename}"
        return [page_from_text(url, filename, text)]

    def _extract(self, data: bytes, ext: str, mime: str) -> str:
        if ext in (".txt", ".md", ".markdown") or mime.startswith("text/"):
            return data.decode("utf-8", errors="replace")

        if ext == ".pdf" or mime == "application/pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            parts = []
            for page in reader.pages:
                parts.append(page.extract_text() or "")
            return "\n\n".join(parts)

        if ext == ".docx" or "wordprocessingml" in mime:
            from docx import Document

            doc = Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

        if ext in (".ppt", ".pptx") or "presentationml" in mime:
            from pptx import Presentation

            prs = Presentation(io.BytesIO(data))
            parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        parts.append(shape.text)
            return "\n".join(parts)

        raise ValueError(f"Unsupported file type: {ext or mime}")
