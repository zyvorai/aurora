"""Render a proposal Artifact as PDF, DOCX, or PPTX for download (Phase 8)."""

from __future__ import annotations

import html
from io import BytesIO
from typing import Any

from gtm_api.models import Artifact

SECTION_ORDER = [
    ("proposal_content", "Proposal"),
    ("sow", "Statement of Work"),
    ("roi_analysis", "ROI Analysis"),
    ("pricing", "Pricing"),
    ("timeline", "Timeline"),
]


def _sections(artifact: Artifact) -> list[tuple[str, str]]:
    meta: dict[str, Any] = artifact.metadata_ or {}
    sections = []
    for key, label in SECTION_ORDER:
        value = meta.get(key)
        if value:
            sections.append((label, str(value)))
    return sections


def _title(artifact: Artifact) -> str:
    meta: dict[str, Any] = artifact.metadata_ or {}
    return str(meta.get("title") or artifact.title or "Proposal")


def render_proposal_pdf(artifact: Artifact) -> bytes:
    from weasyprint import HTML

    title = html.escape(_title(artifact))
    body = "".join(
        f"<h2>{html.escape(label)}</h2><p>{html.escape(text).replace(chr(10), '<br/>')}</p>"
        for label, text in _sections(artifact)
    )
    document = f"""
    <html>
      <head>
        <meta charset="utf-8" />
        <style>
          body {{ font-family: Helvetica, Arial, sans-serif; margin: 2.5cm; color: #1a1a1a; }}
          h1 {{ font-size: 22pt; margin-bottom: 0.2em; }}
          h2 {{ font-size: 14pt; margin-top: 1.4em; margin-bottom: 0.3em; border-bottom: 1px solid #ccc; }}
          p {{ font-size: 10.5pt; line-height: 1.5; white-space: pre-wrap; }}
        </style>
      </head>
      <body>
        <h1>{title}</h1>
        {body}
      </body>
    </html>
    """
    return HTML(string=document).write_pdf()


def render_proposal_docx(artifact: Artifact) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(_title(artifact), level=0)
    for label, text in _sections(artifact):
        document.add_heading(label, level=1)
        for paragraph in text.split("\n"):
            document.add_paragraph(paragraph)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_proposal_pptx(artifact: Artifact) -> bytes:
    from pptx import Presentation
    from pptx.util import Pt

    presentation = Presentation()
    title_layout = presentation.slide_layouts[0]
    body_layout = presentation.slide_layouts[1]

    title_slide = presentation.slides.add_slide(title_layout)
    title_slide.shapes.title.text = _title(artifact)
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = "Generated proposal"

    for label, text in _sections(artifact):
        slide = presentation.slides.add_slide(body_layout)
        slide.shapes.title.text = label
        body = slide.placeholders[1].text_frame
        body.word_wrap = True
        lines = text.split("\n") or [text]
        body.text = lines[0]
        for line in lines[1:]:
            paragraph = body.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(14)

    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


EXPORT_RENDERERS = {
    "pdf": (render_proposal_pdf, "application/pdf"),
    "docx": (
        render_proposal_docx,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ),
    "pptx": (
        render_proposal_pptx,
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ),
}
